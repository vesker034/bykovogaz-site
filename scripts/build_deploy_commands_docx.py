# -*- coding: utf-8 -*-
"""Генерация DOCX со списком команд развёртывания проекта БЫКОВОГАЗ."""
from pathlib import Path

from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_LINE_SPACING

ROOT = Path(__file__).resolve().parent.parent
OUTPUT = ROOT / "docs" / "Развертывание_команды_терминала.docx"

SECTIONS = [
    (
        "1. Подключение к серверу и получение проекта",
        [
            (
                "cd /opt",
                "Переход в каталог, где размещают веб-приложения на Linux-сервере.",
            ),
            (
                "git clone https://github.com/vesker034/BYKOVOGAZ.git",
                "Клонирование репозитория с GitHub на сервер (SberCloud, Ubuntu 22.04).",
            ),
            (
                "cd /opt/BYKOVOGAZ",
                "Переход в корень проекта (manage.py, config/, main/, templates/).",
            ),
        ],
    ),
    (
        "2. Установка системных пакетов (Ubuntu 22.04, Python 3.12)",
        [
            (
                "sudo apt update && sudo apt upgrade -y",
                "Обновление списка пакетов и установка обновлений безопасности ОС.",
            ),
            (
                "sudo apt install -y software-properties-common python3-pip build-essential",
                "PPA-утилиты, pip и компилятор для сборки Pillow и других модулей.",
            ),
            (
                "sudo add-apt-repository -y ppa:deadsnakes/ppa && sudo apt update",
                "Репозиторий deadsnakes: Python 3.12 нужен для Django 6 (на Ubuntu 22.04 по умолчанию Python 3.10).",
            ),
            (
                "sudo apt install -y python3.12 python3.12-venv python3.12-dev",
                "Установка Python 3.12, модуля venv и заголовков для C-расширений.",
            ),
        ],
    ),
    (
        "3. Виртуальное окружение и зависимости Python",
        [
            (
                "cd /opt/BYKOVOGAZ && rm -rf venv",
                "Переход в проект и удаление старого venv (если создавался с Python 3.10/3.11).",
            ),
            (
                "python3.12 -m venv venv",
                "Создание виртуального окружения на Python 3.12.",
            ),
            (
                "source venv/bin/activate",
                "Активация venv; в начале строки появится (venv).",
            ),
            (
                "pip install --upgrade pip && pip install -r requirements.txt",
                "Установка Django 6.0.3, Pillow и остальных зависимостей из requirements.txt.",
            ),
            (
                'python -c "import django; print(django.get_version())"',
                "Проверка версии Django (ожидается 6.0.3).",
            ),
        ],
    ),
    (
        "4. Переменные окружения (.env)",
        [
            (
                "nano /opt/BYKOVOGAZ/.env",
                "Создание файла настроек production (см. раздел 5 и .env.example в репозитории).",
            ),
            (
                'python -c "import secrets; print(secrets.token_urlsafe(50))"',
                "Генерация случайного DJANGO_SECRET_KEY.",
            ),
            (
                "set -a && source .env && set +a",
                "Загрузка переменных в текущую shell-сессию (Django читает os.getenv, не файл .env автоматически).",
            ),
        ],
    ),
    (
        "5. Пример содержимого файла .env",
        [
            (
                "DJANGO_SECRET_KEY=<сгенерированный-ключ>",
                "Секретный ключ Django; тот же ключ укажите в gunicorn.service.",
            ),
            (
                "DJANGO_DEBUG=False",
                "Production-режим. Для временной проверки runserver можно True (тогда dev-сервер отдаёт static/).",
            ),
            (
                "DJANGO_ALLOWED_HOSTS=bykovogazzz.ru,www.bykovogazzz.ru,84.54.59.117,bykovogaz,localhost,127.0.0.1",
                "Разрешённые Host; без домена — ошибка DisallowedHost.",
            ),
            (
                "DJANGO_CSRF_TRUSTED_ORIGINS=https://bykovogazzz.ru,https://www.bykovogazzz.ru",
                "Доверенные Origin для CSRF при HTTPS (вход в /admin/).",
            ),
        ],
    ),
    (
        "6. База данных, админка и статика",
        [
            (
                "python manage.py migrate",
                "Миграции SQLite и начальное наполнение CMS (seed-миграция 0007).",
            ),
            (
                "python manage.py createsuperuser",
                "Учётная запись для /admin/.",
            ),
            (
                "python manage.py collectstatic --noinput",
                "Сбор CSS/JS/изображений в staticfiles/ — их отдаёт Caddy по /static/.",
            ),
            (
                "python manage.py check",
                "Проверка конфигурации Django.",
            ),
        ],
    ),
    (
        "7. Тестовый запуск (опционально, перед production)",
        [
            (
                "python manage.py runserver 0.0.0.0:8000",
                "Dev-сервер для проверки. В браузере: http://84.54.59.117:8000 (не http://0.0.0.0). "
                "При DJANGO_DEBUG=False стили не загрузятся — для production используйте Gunicorn + Caddy.",
            ),
            (
                "curl -I http://127.0.0.1:8000/",
                "Проверка ответа приложения локально на сервере.",
            ),
        ],
    ),
    (
        "8. Gunicorn (systemd, порт 127.0.0.1:8000)",
        [
            (
                "pip install gunicorn",
                "WSGI-сервер для production (в venv, после source venv/bin/activate).",
            ),
            (
                "sudo cp /opt/BYKOVOGAZ/deploy/gunicorn.service.example /etc/systemd/system/gunicorn.service",
                "Копирование unit-файла из репозитория (или создайте вручную через nano).",
            ),
            (
                "sudo nano /etc/systemd/system/gunicorn.service",
                "Подставьте DJANGO_SECRET_KEY из .env; проверьте ALLOWED_HOSTS и CSRF_TRUSTED_ORIGINS.",
            ),
            (
                "sudo systemctl daemon-reload",
                "Перечитывание unit-файлов systemd после изменений.",
            ),
            (
                "sudo systemctl enable gunicorn && sudo systemctl start gunicorn",
                "Автозапуск и старт Gunicorn.",
            ),
            (
                "sudo systemctl status gunicorn",
                "Ожидается: active (running). Выход из просмотра — клавиша q.",
            ),
            (
                "curl -I http://127.0.0.1:8000/",
                "Ожидается HTTP/1.1 200 OK от Gunicorn.",
            ),
        ],
    ),
    (
        "9. Caddy (домен bykovogazzz.ru, HTTPS)",
        [
            (
                "sudo apt install -y debian-keyring debian-archive-keyring apt-transport-https curl",
                "Зависимости для официального репозитория Caddy.",
            ),
            (
                "curl -1sLf 'https://dl.cloudsmith.io/public/caddy/stable/gpg.key' | sudo gpg --dearmor -o /usr/share/keyrings/caddy-stable-archive-keyring.gpg",
                "GPG-ключ репозитория Caddy.",
            ),
            (
                "curl -1sLf 'https://dl.cloudsmith.io/public/caddy/stable/debian.deb.txt' | sudo tee /etc/apt/sources.list.d/caddy-stable.list",
                "Подключение apt-репозитория Caddy.",
            ),
            (
                "sudo apt update",
                "Обновление списка пакетов после добавления репозитория Caddy.",
            ),
            (
                "sudo mkdir -p /etc/caddy",
                "Создание каталога конфигурации (нужно до записи Caddyfile, если Caddy ещё не установлен).",
            ),
            (
                "sudo tee /etc/caddy/Caddyfile > /dev/null << 'EOF' … EOF",
                "Запись конфигурации сайта (см. deploy/Caddyfile в репозитории): домен, /static/, /media/, reverse_proxy на 127.0.0.1:8000.",
            ),
            (
                "cat /etc/caddy/Caddyfile",
                "Проверка, что в файле есть bykovogazzz.ru, а не стандартная заглушка пакета.",
            ),
            (
                "sudo apt install -y caddy",
                "Установка Caddy. При вопросе про Caddyfile нажмите N (оставить свою версию), не Y!",
            ),
            (
                "sudo caddy validate --config /etc/caddy/Caddyfile",
                "Проверка синтаксиса конфигурации.",
            ),
            (
                "sudo systemctl enable caddy && sudo systemctl restart caddy",
                "Автозапуск и применение конфигурации; Let's Encrypt выпускается автоматически при корректном DNS.",
            ),
            (
                "sudo systemctl status caddy",
                "Ожидается: active (running).",
            ),
            (
                "curl -I https://bykovogazzz.ru",
                "Проверка HTTPS с сервера.",
            ),
        ],
    ),
    (
        "10. DNS и firewall (SberCloud)",
        [
            (
                "A @ → 84.54.59.117",
                "У регистратора домена: основная A-запись на публичный IP ВМ.",
            ),
            (
                "A www → 84.54.59.117",
                "A-запись для www.bykovogazzz.ru.",
            ),
            (
                "nslookup bykovogazzz.ru",
                "Проверка DNS (распространение до 24–48 ч).",
            ),
            (
                "Security group: TCP 22, 80, 443",
                "В SberCloud: SSH, HTTP (для ACME/Let's Encrypt), HTTPS.",
            ),
        ],
    ),
    (
        "11. SSH-доступ (Termius)",
        [
            (
                "Подключение: 84.54.59.117, порт 22, пользователь vesker034",
                "После переустановки ОС на сервере; не используйте порт 8080 для SSH.",
            ),
            (
                'ssh-keygen -t rsa -b 4096 -f "$HOME/.ssh/id_rsa_bykovogaz"',
                "На ноутбуке: создание SSH-ключа для входа без пароля.",
            ),
            (
                "mkdir -p ~/.ssh && chmod 700 ~/.ssh && nano ~/.ssh/authorized_keys",
                "На сервере: добавление публичного ключа (.pub).",
            ),
            (
                "chmod 600 ~/.ssh/authorized_keys",
                "Права на authorized_keys.",
            ),
        ],
    ),
    (
        "12. Обслуживание и диагностика",
        [
            (
                "git pull && source venv/bin/activate && set -a && source .env && set +a",
                "Обновление кода и переменных окружения.",
            ),
            (
                "python manage.py migrate && python manage.py collectstatic --noinput",
                "После обновления кода: миграции и пересбор статики.",
            ),
            (
                "sudo systemctl restart gunicorn && sudo systemctl restart caddy",
                "Перезапуск приложения и веб-сервера.",
            ),
            (
                "journalctl -u gunicorn -n 30 --no-pager",
                "Логи Gunicorn при ошибках.",
            ),
            (
                "journalctl -u caddy -n 30 --no-pager",
                "Логи Caddy (HTTPS, сертификаты, прокси).",
            ),
            (
                "sudo tee /etc/caddy/Caddyfile … && sudo systemctl restart caddy",
                "Если при apt install caddy случайно нажали Y и перезаписали Caddyfile — восстановите конфиг и перезапустите.",
            ),
        ],
    ),
]

ARCHITECTURE_NOTE = (
    "Схема production: браузер → Caddy (порты 80/443, HTTPS, /static/, /media/) → "
    "Gunicorn (127.0.0.1:8000) → Django → SQLite (db.sqlite3). "
    "Публичный адрес: https://bykovogazzz.ru , админка: https://bykovogazzz.ru/admin/ ."
)


def _set_normal_style(doc: Document) -> None:
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(14)
    pf = style.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = 1.5
    pf.space_after = Pt(6)


def build() -> Path:
    doc = Document()
    _set_normal_style(doc)

    sections = doc.sections[0]
    sections.top_margin = Cm(2)
    sections.bottom_margin = Cm(2)
    sections.left_margin = Cm(3)
    sections.right_margin = Cm(1.5)

    title = doc.add_paragraph()
    title.alignment = 1
    run = title.add_run("Команды терминала для развёртывания сайта БЫКОВОГАЗ")
    run.bold = True
    run.font.size = Pt(16)

    doc.add_paragraph(
        "Документ описывает развёртывание Django-проекта на виртуальной машине Ubuntu 22.04 "
        "(SberCloud): IP 84.54.59.117, домен bykovogazzz.ru, каталог /opt/BYKOVOGAZ. "
        "Production-стек: Gunicorn + Caddy + SQLite."
    )
    doc.add_paragraph(ARCHITECTURE_NOTE)

    for section_title, items in SECTIONS:
        h = doc.add_paragraph()
        h_run = h.add_run(section_title)
        h_run.bold = True
        h_run.font.size = Pt(14)

        for cmd, explanation in items:
            p_cmd = doc.add_paragraph()
            p_cmd.paragraph_format.left_indent = Cm(0.5)
            r_cmd = p_cmd.add_run("Команда: ")
            r_cmd.bold = True
            p_cmd.add_run(cmd)

            p_exp = doc.add_paragraph()
            p_exp.paragraph_format.left_indent = Cm(1)
            r_exp = p_exp.add_run("Назначение: ")
            r_exp.italic = True
            p_exp.add_run(explanation)

    doc.add_paragraph()
    note = doc.add_paragraph()
    note_run = note.add_run("Примечание. ")
    note_run.bold = True
    note.add_run(
        "Перед командами Django выполняйте: cd /opt/BYKOVOGAZ && source venv/bin/activate "
        "&& set -a && source .env && set +a. "
        "Файлы deploy/Caddyfile и deploy/gunicorn.service.example — эталоны конфигурации в репозитории."
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    path = build()
    print(f"Создан файл: {path}")
